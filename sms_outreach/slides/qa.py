"""Created by JXP and Claude.

QA checks for a built Kraw-style deck.

Failures (always):
  1. main-deck slide count > max_slides (backup slides excluded);
  2. margins: a shape we placed outside the side margins or below
     layout.BOTTOM_LIMIT (Kraw's own title box, the slide number and the
     title slide are exempt);
  3. fonts: an explicit typeface other than Roboto (text with no typeface
     inherits the master);
  4. a title that needs more than two lines;
  5. a sub-line that does not fit on one line (it would run into the source
     line and the logo);
  6. a bullet list taller than its box.
Warnings (failures with strict=True, e.g. for the final deck):
  7. a slide without speaker notes;
  8. a figure (picture) without a source line.
Placeholders (red boxes for missing figures) are counted and listed.
"""
from pptx import Presentation
from pptx.util import Emu

from . import layout as L

PLACED = {"Figure", "Caption", "Bullets", "Source", "PLACEHOLDER"}


def inches(v):
    """Created by JXP and Claude. EMU -> inches."""
    return Emu(v).inches


def shape_box(sh):
    """Created by JXP and Claude. (x, y, w, h) of a shape in inches."""
    return inches(sh.left), inches(sh.top), inches(sh.width), inches(sh.height)


def run_size(sh, default):
    """Created by JXP and Claude. Font size (pt) of the first sized run in a text shape."""
    for para in sh.text_frame.paragraphs:
        for run in para.runs:
            if run.font.size:
                return run.font.size.pt
    return default


def placeholders(prs):
    """Created by JXP and Claude. [(slide number, title, label)] for every
    PLACEHOLDER box in the deck."""
    out = []
    for i, slide in enumerate(prs.slides, start=1):
        title = next((sh.text_frame.text for sh in slide.shapes if sh.name == "Title"), "")
        for sh in slide.shapes:
            if sh.name == "PLACEHOLDER":
                out.append((i, title, sh.text_frame.text.split("\n", 1)[-1]))
    return out


def check(path, max_slides=None, n_backup=0, strict=False):
    """Created by JXP and Claude. Run all checks.

    Inputs: deck path; max_slides (None = no limit); n_backup = number of
    trailing backup slides not counted against max_slides; strict = treat
    warnings as failures.
    Output: (problems, warnings, placeholder list).
    """
    prs = Presentation(str(path))
    problems, warnings = [], []
    n = len(prs.slides)
    if max_slides and n - n_backup > max_slides:
        problems.append(f"main deck has {n - n_backup} slides (> {max_slides})")
    for i, slide in enumerate(prs.slides, start=1):
        shapes = {sh.name: sh for sh in slide.shapes}
        title = shapes["Title"].text_frame.text if "Title" in shapes else ""
        tag = f"slide {i} '{title[:30]}'"
        for sh in slide.shapes:
            if i > 1 and sh.name in PLACED and not L.inside(shape_box(sh)):
                x, y, w, h = shape_box(sh)
                problems.append(f"{tag}: {sh.name} outside margins (l={x:.2f} r={x + w:.2f} b={y + h:.2f})")
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    for run in para.runs:
                        if run.font.name and run.font.name != "Roboto":
                            problems.append(f"slide {i}: font {run.font.name!r} in {sh.name}")
        if title and L.wrapped_lines(title, L.title_size(title)[0], L.TITLE_TEXT_W) > 2:
            problems.append(f"{tag}: title needs more than two lines")
        if "Sub-line" in shapes:
            sub = shapes["Sub-line"]
            if L.text_width(sub.text_frame.text, run_size(sub, L.SUBLINE_PT)) > L.SUBLINE_TEXT_W:
                problems.append(f"{tag}: sub-line wraps (shorten it)")
        if "Bullets" in shapes:
            bul = shapes["Bullets"]
            _, _, w, h = shape_box(bul)
            lines = [p.text for p in bul.text_frame.paragraphs]
            if L.bullets_height(lines, run_size(bul, L.BULLET_PT), w) > h + 0.01:
                problems.append(f"{tag}: bullets overflow their box (cut text)")
        if i > 1 and not (slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip()):
            warnings.append(f"{tag}: no speaker notes")
        if "Figure" in shapes and "Source" not in shapes:
            warnings.append(f"{tag}: figure without a source line")
    if strict:
        problems, warnings = problems + warnings, []
    return problems, warnings, placeholders(prs)


def report(path, max_slides=None, n_backup=0, strict=False):
    """Created by JXP and Claude. Print the QA report; return True if it passed."""
    problems, warnings, ph = check(path, max_slides, n_backup, strict)
    print(f"QA {path}: {len(ph)} placeholder(s), {len(warnings)} warning(s)")
    for i, title, label in ph:
        print(f"  placeholder  slide {i:2d} {title[:40]:40s} {label.replace(chr(10), ' ')[:60]}")
    for w in warnings:
        print("  warn", w)
    for p in problems:
        print("  FAIL", p)
    print("  ALL CHECKS PASSED" if not problems else f"  {len(problems)} problem(s)")
    return not problems
