"""Created by JXP and Claude.

Gather every existing image the WMKO 2026 plan needs into
talks/wmko_2026/figures/, and record each one's origin, credit and license
in figures/sources.json and figures/credits.md.

Each entry of ASSETS maps a target file name (as used in slides/plan.py) to
one source:
    File(path)                          copy a file
    Picture(pptx, slide, pick)          extract an embedded picture from a deck
                                        (pick = index among the slide's pictures in
                                        shape order, or "largest")
    SlideImage(pptx, slide)             render a whole slide (LibreOffice -> PDF -> PNG)
    Url(url)                            download
and optional `crop` = (left, top, right, bottom) as fractions of the image.

Usage (ocean14 env, from the repo root):
    python talks/wmko_2026/scripts/harvest_assets.py [--only NAME ...] [--force]
Existing targets are skipped unless --force.  Rendered decks are cached in
figures/.cache/ (git-ignored).
"""
import argparse
import io
import json
import shutil
import subprocess
import urllib.request
from dataclasses import dataclass
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

TALK = Path(__file__).resolve().parents[1]
REPO = TALK.parents[1]
FIGS = TALK / "figures"
CACHE = FIGS / ".cache"
HOME = Path.home()

KRAW = HOME / "Projects/ClimateIntelligence/presentations/2026_WMKO/Kraw_2024.pptx"
MBARI = REPO / "context/MBARI_2026.pptx"
HONO = REPO / "context/Honokaʻa_HS_2026.pptx"
HINZ = REPO / "context/Integrating Claude into instrumentation research.pptx"
CI_FIGS = HOME / "Projects/ClimateIntelligence/presentations/2026_WMKO/figs"
EXO = HOME / "Projects/PypeIt/first-hires-exoplanet"
UA = "Mozilla/5.0 (sms-outreach harvest_assets; jxp@ucsc.edu)"
MAX_SIDE = 2400                # stored images are capped at this many pixels


@dataclass
class File:
    path: Path


@dataclass
class Picture:
    pptx: Path
    slide: int
    pick: object = "largest"


@dataclass
class SlideImage:
    pptx: Path
    slide: int
    dpi: int = 200


@dataclass
class Url:
    url: str


@dataclass
class Asset:
    """Created by JXP and Claude. One target image and where it comes from."""
    src: object
    source: str             # what it is / where it is from (for notes and credits.md)
    credit: str = ""
    license: str = ""
    crop: tuple = None      # (left, top, right, bottom) fractions
    circle: bool = False    # mask to the inscribed circle (transparent corners; target must be .png)


# ----------------------------------------------------------------- ASSETS
# Filled from the B2 survey of each deck / repo / web page (see the log).
# "check" in a credit/license marks an origin we could not confirm.
UNKNOWN = "unknown — check"
OWN = "own work"
ARXIV = "arXiv.org perpetual non-exclusive license (authors retain copyright); reused with attribution by a co-author"

ASSETS = {
    # --- Kraw 2024 lecture -------------------------------------------------------------
    "kraw_sea_unknown.jpg": Asset(Picture(KRAW, 2, 0), "Sunset over the Pacific (Kraw 2024 slide 2, left)",
                                  credit="J. X. Prochaska? — check", license=UNKNOWN,
                                  crop=(0.38, 0.46, 1.0, 1.0)),
    "kraw_sky_unknown.jpg": Asset(Picture(KRAW, 2, 1), "JWST's First Deep Field, SMACS 0723 (Kraw 2024 slide 2, right)",
                                  credit="NASA, ESA, CSA, STScI", license="public domain (NASA)",
                                  crop=(0.0, 0.0, 0.60, 1.0)),
    "kraw_cats_dogs.jpg": Asset(Picture(KRAW, 14, 0), "AI 2012: three dogs and two cats (Kraw 2024 slide 14)",
                                credit="stock photo — check", license=UNKNOWN),
    "kraw_parks_cnn.png": Asset(Picture(KRAW, 19, 1),
                                "Lick quasar spectrum Q2138-4427 with a damped Lyα absorber, the kind of data "
                                "Parks et al. 2018 fed to a CNN (Kraw 2024 slide 19)",
                                credit="J. X. Prochaska (Lick 3-m data)", license=OWN),
    "ulmo_outliers.png": Asset(Picture(KRAW, 31, 1), "Ulmo: nine SST outlier cutouts (Kraw 2024 slide 31)",
                               credit="Prochaska, Cornillon & Reiman 2021 (Ulmo)", license=OWN,
                               crop=(0.115, 0.11, 0.905, 0.895)),
    # --- MBARI 2026 talk ---------------------------------------------------------------
    "sst_cutout.png": Asset(Picture(MBARI, 14, 2), "VIIRS SST cutout off California (MBARI 2026 slide 14)",
                            credit="NOAA VIIRS SST; plot via J. X. Prochaska, MBARI 2026", license="NOAA data public domain",
                            crop=(0.10, 0.0, 0.80, 0.92)),
    "cmb_sky.png": Asset(SlideImage(MBARI, 14, dpi=300), "Planck CMB map over a starry night sky (MBARI 2026 slide 14, rendered)",
                         credit="ESA and the Planck Collaboration; night-sky photo — check",
                         license="ESA Standard Licence (credit required); sky photo unknown — check",
                         crop=(0.008, 0.169, 0.474, 0.926)),
    "mbari_cnn.png": Asset(Picture(MBARI, 3, 1), "VGG-16 convolutional neural network diagram (MBARI 2026 slide 3)",
                           credit="after D. Frossard 2016, 'VGG in TensorFlow' — check", license=UNKNOWN),
    "nenya_umap.png": Asset(Picture(MBARI, 24, 0), "Nenya: UMAP of the VIIRS SST manifold (MBARI 2026 slide 24)",
                            credit="J. X. Prochaska et al. (Nenya)", license=OWN, crop=(0.03, 0.0, 0.70, 0.95)),
    "enki_reconstruction.png": Asset(Picture(MBARI, 33, 0),
                                     "Enki: original / masked / reconstructed / residual SST (MBARI 2026 slide 33)",
                                     credit="Agabin, Prochaska et al. 2024 (Enki)", license=OWN),
    "ulmo_alt.png": Asset(Picture(MBARI, 5, 0), "Ulmo: SST outlier cutouts (MBARI 2026 slide 5)",
                          credit="J. X. Prochaska et al. (Ulmo)", license=OWN, crop=(0.11, 0.10, 0.91, 0.91)),
    "launch_art.png": Asset(Picture(MBARI, 6, 0), "Butterflies forming a '5': Claude launch art (MBARI 2026 slide 6)",
                            credit="Anthropic — check", license="Anthropic copyright; editorial use in a talk — check"),
    "pab_infographic.png": Asset(Picture(MBARI, 40, 0), "PAB: PACE satellite x BGC-Argo floats (MBARI 2026 slide 40)",
                                 credit="J. X. Prochaska & Claude (PAB)", license=OWN),
    "team_ai_remote_sensing.png": Asset(SlideImage(MBARI, 4), "'AI on Remote Sensing' team (MBARI 2026 slide 4, rendered)",
                                        credit="team members' photos", license="as used in MBARI 2026",
                                        crop=(0.0, 0.095, 1.0, 1.0)),
    # --- Honokaʻa High School 2026 -------------------------------------------------------
    "login_card.jpg": Asset(Picture(HONO, 2, 0), "Keck Observer Portal 'Your Information' card with a 1990s headshot "
                                                 "(Honokaʻa 2026 slide 2)",
                            credit="J. X. Prochaska (screenshot)", license=OWN, crop=(0.0, 0.23, 1.0, 0.86)),
    "headshot_1990s.png": Asset(Picture(HONO, 2, 0), "J. X. Prochaska in the 1990s: the photo on his Keck Observer "
                                                     "Portal card (Honokaʻa 2026 slide 2)",
                                credit="J. X. Prochaska", license=OWN, circle=True,
                                crop=(0.4980, 0.2593, 0.7510, 0.4175)),
    "keck_domes.jpg": Asset(Picture(HONO, 3, 0), "Keck I and II domes on Maunakea (Honokaʻa 2026 slide 3)",
                            credit="W. M. Keck Observatory? — check", license=UNKNOWN),
    # --- rendered slides ---------------------------------------------------------------
    "hinz_slumping.png": Asset(SlideImage(HINZ, 2), "Phil Hinz: 'Developing a viscous model for glass slumping' (KASM)",
                               credit="Phil Hinz (UCSC)", license="used with permission"),
    "hinz_plots.png": Asset(SlideImage(HINZ, 2), "Phil Hinz: sand-height and deflection plots from the glass-slumping "
                                                 "model (KASM)",
                            credit="Phil Hinz (UCSC)", license="used with permission", crop=(0.565, 0.235, 0.935, 0.73)),
    "exo_public_1.png": Asset(SlideImage(EXO / "docs/slides/public_summary.pptx", 1),
                              "'You could do most of this yourself' (first-hires-exoplanet)",
                              credit="J. X. Prochaska & Claude", license=OWN, crop=(0.0, 0.14, 1.0, 0.83)),
    "exo_public_2.png": Asset(SlideImage(EXO / "docs/slides/public_summary.pptx", 2),
                              "'A teacher could lead a class through it' (first-hires-exoplanet)",
                              credit="J. X. Prochaska & Claude", license=OWN),
    # --- files from other repos ----------------------------------------------------------
    "sacbee.png": Asset(File(HOME / "Projects/ClimateIntelligence/presentations/2026_WMKO/sacbee.png"),
                        "'California professor: AI surpassed me as a scientist. What now?', Sacramento Bee, 21 Aug 2026",
                        credit="J. X. Prochaska, The Sacramento Bee",
                        license="author's op-ed (screenshot, cropped to the headline)", crop=(0.0, 0.0, 1.0, 0.355)),
    "a2_bio_uplift.png": Asset(File(CI_FIGS / "a2_bio_uplift.png"), "AI uplift on biology tasks, 2024–2026",
                               credit="JXP & Claude; data: RAND 2024, OpenAI 2024, Anthropic 2025, Zhang+ 2026, "
                                      "Götting+ 2025", license=OWN),
    "a3_hacking.png": Asset(File(CI_FIGS / "a3_hacking.png"), "Claude Mythos Preview and the Hugging Face break-in",
                            credit="JXP & Claude; Scientific American & Axios, Apr 2026; CNBC, Jul 2026", license=OWN),
    "a5_arxiv.png": Asset(File(CI_FIGS / "a5_arxiv.png"), "New arXiv submissions per month, 1991–2026",
                          credit="JXP & Claude; data: arXiv stats, arXiv blog 1 Oct 2026", license=OWN),
    "a6_proposals.png": Asset(File(CI_FIGS / "a6_proposals.png"), "Fifteen 'Excellent' proposals, one funded",
                              credit="JXP & Claude", license=OWN),
    "s4_senses.png": Asset(File(CI_FIGS / "s4_senses.png"), "Our senses: eyes, ears, nose, skin",
                           credit="JXP & Claude; Hecht+ 1942, USGS, Skedung+ 2013", license=OWN),
    "hires.jpg": Asset(File(CI_FIGS / "images/hires.jpg"), "HIRES spectrometer, Keck I",
                       credit="Courtesy W. M. Keck Observatory",
                       license="WMKO permission for this talk (Q8)"),
    "keck_primary.jpg": Asset(File(CI_FIGS / "images/keck_primary.jpg"), "Keck segmented primary mirror",
                              credit="z2amiller / Wikimedia Commons", license="CC BY-SA 2.0"),
    "loc_over_time.png": Asset(File(HOME / "bin/reports/figs/loc_over_time.png"),
                               "Lines of code written by J. X. Prochaska, 1995–2026",
                               credit="JXP & Claude (~/bin/reports/x_lines_of_code.md)", license=OWN,
                               crop=(0.0, 0.0, 1.0, 0.61)),
    "fig7_three_ways.png": Asset(File(EXO / "docs/figs/fig7_three_ways.png"),
                                 "HD 187123 b three ways: lamp, iodine + open software, modern pipeline",
                                 credit="JXP & Claude; right panel data Teklu et al. 2025", license=OWN),
    "holy_grail_arc.png": Asset(File(HOME / "Projects/PypeIt/the-holy-grail/PR/pr_figure_2panels.png"),
                                "A thorium-argon arc before and after blind wavelength calibration",
                                credit="JXP & Claude; APF/Lick Observatory data; PypeIt", license=OWN),
    "holy_grail_raw.png": Asset(File(HOME / "Projects/PypeIt/the-holy-grail/PR/pr_figure_2panels.png"),
                                "A raw thorium-argon arc exposure (APF, Lick)", credit="JXP & Claude; APF/Lick Observatory data",
                                license=OWN, crop=(0.0, 0.0, 0.49, 1.0)),
    "holy_grail_cal.png": Asset(File(HOME / "Projects/PypeIt/the-holy-grail/PR/pr_figure_2panels.png"),
                                "The same arc coloured by the wavelength found with no human input",
                                credit="JXP & Claude; PypeIt", license=OWN, crop=(0.51, 0.0, 1.0, 1.0)),
    "fig01_inverse_problem.png": Asset(File(HOME / "Projects/claudes-phd-thesis/reports/figures/fig01_inverse_problem.png"),
                                       "Claude's PhD: one ocean-colour spectrum in, five constituent spectra out",
                                       credit="Claude (candidate) & JXP; Loisel et al. 2023 synthetic data", license=OWN,
                                       crop=(0.0, 0.06, 1.0, 0.50)),
    "ai_da_timeline.png": Asset(File(HOME / "Projects/BOONUS/reports/data_assimilation/figs/ai_da_timeline.png"),
                                "AI in data assimilation, 2018–2026: weather vs ocean",
                                credit="JXP & Claude (BOONUS)", license=OWN),
    "llc4320_sst.png": Asset(File(HOME / "Oceanography/python/wrangler/docs/slides/figs/llc4320_v2_sst_gulfstream.png"),
                             "LLC4320 virtual ocean: Gulf Stream sea-surface temperature at ~2 km",
                             credit="JXP & Claude (wrangler); LLC4320 MITgcm, NASA/JPL ECCO (Menemenlis et al.)",
                             license="own figure; NASA data"),
    # --- downloads -----------------------------------------------------------------------
    "wolfe_disk.jpg": Asset(Url("https://public.nrao.edu/wp-content/uploads/2020/04/"
                                "nrao20in04_WolfeDisk_illustration_SD.jpg"),
                            "Artist's impression of the Wolfe Disk (Neeleman et al. 2020, Nature)",
                            credit="NRAO/AUI/NSF, S. Dagnello", license="CC BY 4.0 (NRAO media resources)"),
    "frb_jwst_nircam.png": Asset(Url("https://arxiv.org/html/2508.01648v1/nircam.png"),
                                 "JWST/NIRCam view of the host of FRB 20240304B, z = 2.148 (Caleb et al., Fig. 2A-B)",
                                 credit="Caleb et al. 2025 (arXiv:2508.01648); NASA, ESA, CSA, JWST", license=ARXIV,
                                 crop=(0.0, 0.0, 1.0, 0.38)),
    "frb_keck_lris.png": Asset(Url("https://arxiv.org/html/2508.01648v1/FRB20240304_LRIS_R-band.png"),
                               "Keck/LRIS R-band image at the FRB 20240304B position: no host (Caleb et al., Fig. S5)",
                               credit="Caleb et al. 2025; Keck I/LRIS program U299 (PI Prochaska)", license=ARXIV),
    "hokulea_1976.jpg": Asset(Url("https://upload.wikimedia.org/wikipedia/commons/d/d8/Hokule%27a.jpg"),
                              "Hōkūleʻa arriving in Honolulu from Tahiti, 1976",
                              credit="Phil Uhl / Wikimedia Commons", license="CC BY-SA 3.0"),
    "frb_lris_zoom.png": Asset(Url("https://arxiv.org/html/2508.01648v1/FRB20240304_LRIS_R-band.png"),
                               "Keck/LRIS R-band zoom on the FRB 20240304B position: nothing there "
                               "(Caleb et al., Fig. S5, right panel)",
                               credit="Caleb et al. 2025; Keck I/LRIS program U299 (PI Prochaska)", license=ARXIV,
                               crop=(0.592, 0.131, 0.952, 0.765)),
    "frb_jwst_zoom.png": Asset(Url("https://arxiv.org/html/2508.01648v1/nircam.png"),
                               "JWST/NIRCam zoom on the FRB 20240304B position: a faint host galaxy at z = 2.148 "
                               "(Caleb et al., Fig. 2B)",
                               credit="Caleb et al. 2025 (arXiv:2508.01648); NASA, ESA, CSA, JWST", license=ARXIV,
                               crop=(0.616, 0.042, 0.974, 0.329)),
    "lgs_laser.jpg": Asset(Url("https://upload.wikimedia.org/wikipedia/commons/9/9b/"
                               "The_Stars_Above_Maunakea_%28_MG_7892-Panorama2-CC%29.jpg"),
                           "Laser guide stars from the two Keck telescopes on Maunakea (NOIRLab panorama)",
                           credit="International Gemini Observatory/NOIRLab/NSF/AURA/B. Tafreshi / Wikimedia Commons",
                           license="CC BY 4.0", crop=(0.0, 0.0, 0.60, 0.80)),
    "iodine_cell.jpg": Asset(Url("https://upload.wikimedia.org/wikipedia/commons/thumb/e/e6/I_vapor.png/"
                                 "1920px-I_vapor.png"),
                             "Purple iodine vapour in a glass vessel (stand-in for an astronomical iodine cell)",
                             credit="2x910 / Wikimedia Commons", license="CC BY-SA 4.0", crop=(0.08, 0.0, 0.92, 1.0)),
    "alma.jpg": Asset(Url("https://upload.wikimedia.org/wikipedia/commons/7/71/"
                          "ALMA_beneath_the_stars_%28duro_4776-cc%29.jpg"),
                      "ALMA antennas on the Chajnantor plateau beneath the Milky Way",
                      credit="A. Duro/ESO / Wikimedia Commons", license="CC BY 4.0"),
    "ocean_glider.jpg": Asset(Url("https://upload.wikimedia.org/wikipedia/commons/c/c5/"
                                  "Global_Profiling_Glider_Deployment_%28OOI_111%29.jpg"),
                              "An ocean glider (OOI Slocum) at the surface before its first dive (stand-in for a Spray)",
                              credit="Ocean Observatories Initiative / Wikimedia Commons", license="public domain"),
    "lick_observatory.jpg": Asset(Url("https://upload.wikimedia.org/wikipedia/commons/thumb/5/58/"
                                      "Mt_Hamilton_and_Lick_Observatory_%285265825018%29.jpg/"
                                      "1920px-Mt_Hamilton_and_Lick_Observatory_%285265825018%29.jpg"),
                                  "Lick Observatory on Mt Hamilton after a snowfall, 15 Dec 2010",
                                  credit="Jitze Couperus / Wikimedia Commons", license="CC BY 2.0"),
}


# ---------------------------------------------------------------- helpers
def pictures(slide):
    """Created by JXP and Claude. Picture shapes on a slide in shape order
    (descending into groups)."""
    out = []

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(sh.shapes)
            elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(sh, "image"):
                try:
                    sh.image
                except (AttributeError, ValueError):
                    continue
                out.append(sh)
    walk(slide.shapes)
    return out


_DECKS = {}


def open_deck(path):
    """Created by JXP and Claude. python-pptx Presentation, cached per path."""
    if path not in _DECKS:
        _DECKS[path] = Presentation(str(path))
    return _DECKS[path]


def extract_picture(src):
    """Created by JXP and Claude. (bytes, extension) of an embedded picture."""
    slide = open_deck(src.pptx).slides[src.slide - 1]
    pics = pictures(slide)
    if not pics:
        raise ValueError(f"{src.pptx.name} slide {src.slide}: no pictures")
    if src.pick == "largest":
        pic = max(pics, key=lambda p: p.image.size[0] * p.image.size[1])
    else:
        pic = pics[src.pick]
    return pic.image.blob, pic.image.ext


def render_slide(src):
    """Created by JXP and Claude. PNG bytes of one slide rendered with
    LibreOffice (the deck's PDF is cached)."""
    CACHE.mkdir(parents=True, exist_ok=True)
    pdf = CACHE / (src.pptx.stem + ".pdf")
    if not pdf.exists():
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(CACHE), str(src.pptx)],
                       check=True, capture_output=True)
    # LibreOffice skips hidden slides, so map the slide number to its PDF page
    slides = list(open_deck(src.pptx).slides)
    if slides[src.slide - 1]._element.get("show") == "0":
        raise ValueError(f"{src.pptx.name} slide {src.slide} is hidden (not in the PDF)")
    page = sum(s._element.get("show") != "0" for s in slides[:src.slide])
    stem = CACHE / f"{src.pptx.stem}_s{src.slide}"
    subprocess.run(["pdftoppm", "-r", str(src.dpi), "-f", str(page), "-l", str(page), "-png",
                    "-singlefile", str(pdf), str(stem)], check=True)
    return Path(str(stem) + ".png").read_bytes(), "png"


def download(src):
    """Created by JXP and Claude. Bytes of a URL (with a polite User-Agent)."""
    req = urllib.request.Request(src.url, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=60) as r:
        data = r.read()
    return data, src.url.rsplit(".", 1)[-1].split("?")[0].lower()


def fetch(asset):
    """Created by JXP and Claude. (bytes, extension) for an asset's source."""
    src = asset.src
    if isinstance(src, File):
        return Path(src.path).read_bytes(), Path(src.path).suffix.lstrip(".").lower()
    if isinstance(src, Picture):
        return extract_picture(src)
    if isinstance(src, SlideImage):
        return render_slide(src)
    if isinstance(src, Url):
        return download(src)
    raise TypeError(src)


def write_image(name, data, crop, circle=False):
    """Created by JXP and Claude. Save `data` as figures/<name>, converting
    the format to match the target extension, applying `crop` and, if
    `circle`, a circular transparency mask."""
    out = FIGS / name
    im = Image.open(io.BytesIO(data))
    if crop:
        l, t, r, b = crop
        w, h = im.size
        im = im.crop((round(l * w), round(t * h), round(r * w), round(b * h)))
    im.thumbnail((MAX_SIDE, MAX_SIDE))     # keep the repo small; slides need far less
    if circle:
        im = im.convert("RGBA")
        mask = Image.new("L", im.size, 0)
        ImageDraw.Draw(mask).ellipse((0, 0, im.size[0] - 1, im.size[1] - 1), fill=255)
        im.putalpha(mask)
    if out.suffix.lower() in (".jpg", ".jpeg"):
        im.convert("RGB").save(out, quality=92)
    else:
        if im.mode == "CMYK":
            im = im.convert("RGB")
        im.save(out)
    return out, im.size


def describe(src):
    """Created by JXP and Claude. Short human-readable origin of a source."""
    if isinstance(src, File):
        return str(Path(src.path)).replace(str(HOME), "~")
    if isinstance(src, Picture):
        return f"{src.pptx.name}, slide {src.slide}, picture {src.pick}"
    if isinstance(src, SlideImage):
        return f"{src.pptx.name}, slide {src.slide} (rendered)"
    return src.url


def write_manifests():
    """Created by JXP and Claude. Update sources.json (keeping entries made
    by figure scripts) and rewrite credits.md from ASSETS."""
    manifest_path = FIGS / "sources.json"
    manifest = json.loads(manifest_path.read_text()) if manifest_path.exists() else {}
    for name, a in ASSETS.items():
        entry = manifest.get(name, {})
        entry.update({"origin": describe(a.src), "source": a.source, "credit": a.credit, "license": a.license})
        manifest[name] = entry
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = ["# Image credits — WMKO 2026 \"Sea meets the stars\"", "",
             "Generated by `talks/wmko_2026/scripts/harvest_assets.py`; figures made by our own scripts "
             "are listed in `sources.json` with origin `generated`.", "",
             "| File | What | Credit | License | Origin |", "|---|---|---|---|---|"]
    for name, a in sorted(ASSETS.items()):
        cells = [name, a.source, a.credit or "—", a.license or "—", describe(a.src)]
        lines.append("| " + " | ".join(c.replace("|", "/") for c in cells) + " |")
    (FIGS / "credits.md").write_text("\n".join(lines) + "\n")


def main():
    """Created by JXP and Claude. Harvest all (or --only) assets."""
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--only", nargs="+", help="target names to (re)harvest")
    p.add_argument("--force", action="store_true", help="overwrite existing targets")
    args = p.parse_args()
    FIGS.mkdir(parents=True, exist_ok=True)
    names = args.only or sorted(ASSETS)
    failed = []
    for name in names:
        out = FIGS / name
        if out.exists() and not args.force:
            print(f"skip   {name} (exists)")
            continue
        try:
            data, _ = fetch(ASSETS[name])
            out, size = write_image(name, data, ASSETS[name].crop, ASSETS[name].circle)
            print(f"wrote  {name:28s} {size[0]}x{size[1]}  <- {describe(ASSETS[name].src)}")
        except Exception as err:  # keep going; report at the end
            failed.append(name)
            print(f"FAILED {name}: {err}")
    write_manifests()
    print(f"{len(names) - len(failed)} ok, {len(failed)} failed {failed if failed else ''}")
    shutil.rmtree(CACHE / "tmp", ignore_errors=True)


if __name__ == "__main__":
    main()
